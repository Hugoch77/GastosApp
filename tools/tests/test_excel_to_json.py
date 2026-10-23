"""Pruebas del conversor Excel → respaldo JSON."""

import json
from datetime import date, datetime, timezone

import pytest
from openpyxl import Workbook

from excel_to_json import (
    CardRules,
    Row,
    build_backup,
    guess_category,
    installment_schedule,
    main,
    parse_detail,
    period_for_cut,
    read_rows,
    split_installments,
    statement_cut_for,
    to_cents,
)

# Corte 5, pago 25; lo del día de corte entra al estado siguiente.
BBVA = CardRules(cut_day=5, due_day=25, cut_day_in_next_period=True)
SAME_DAY = CardRules(cut_day=15, due_day=5, cut_day_in_next_period=False)
NOW = datetime(2026, 10, 23, 18, 0, tzinfo=timezone.utc)
TODAY = date(2026, 10, 23)  # facturado: 05/09–04/10 (corte 05/10); actual: 05/10–04/11


def row(description, purchase, amount_cents, months=1, number=1):
    return Row(number, description, purchase, amount_cents, months, None)


def ids():
    counter = iter(range(1, 10_000))
    return lambda: f"id-{next(counter)}"


def build(rows, **overrides):
    options = dict(
        card_name="BBVA Platinum",
        rules=BBVA,
        payment_buffer_days=2,
        today=TODAY,
        billed_paid_on=None,
        now=NOW,
        new_id=ids(),
    )
    options.update(overrides)
    return build_backup(rows, **options)


class TestDates:
    @pytest.mark.parametrize(
        "purchase, expected_cut",
        [
            (date(2026, 10, 4), date(2026, 10, 5)),
            (date(2026, 10, 5), date(2026, 11, 5)),  # día de corte → siguiente
            (date(2026, 12, 20), date(2027, 1, 5)),  # cruza de año
        ],
    )
    def test_statement_cut_next_period_rule(self, purchase, expected_cut):
        assert statement_cut_for(purchase, BBVA) == expected_cut

    def test_statement_cut_same_day_rule(self):
        assert statement_cut_for(date(2026, 10, 15), SAME_DAY) == date(2026, 10, 15)

    def test_cut_day_31_is_clamped(self):
        rules = CardRules(cut_day=31, due_day=20, cut_day_in_next_period=False)
        assert statement_cut_for(date(2027, 2, 10), rules) == date(2027, 2, 28)

    def test_periods(self):
        assert period_for_cut(date(2026, 10, 5), BBVA) == (date(2026, 9, 5), date(2026, 10, 4))
        assert period_for_cut(date(2026, 10, 15), SAME_DAY) == (date(2026, 9, 16), date(2026, 10, 15))

    def test_split_installments_is_exact(self):
        assert split_installments(922040, 6) == [153674, 153674, 153673, 153673, 153673, 153673]
        assert sum(split_installments(100, 3)) == 100

    def test_installment_schedule(self):
        schedule = installment_schedule(date(2026, 9, 30), 90000, 3, BBVA)
        assert schedule == [
            (date(2026, 10, 5), 30000),
            (date(2026, 11, 5), 30000),
            (date(2026, 12, 5), 30000),
        ]


class TestParsing:
    @pytest.mark.parametrize(
        "detail, expected",
        [
            ("Tacos 8/10/26", ("Tacos", date(2026, 10, 8))),
            ("Uber Eats 30/9/2026", ("Uber Eats", date(2026, 9, 30))),
            ("  Gasolina 05/01/27 ", ("Gasolina", date(2027, 1, 5))),
            ("1/10/26", ("Compra", date(2026, 10, 1))),
        ],
    )
    def test_parse_detail(self, detail, expected):
        assert parse_detail(detail) == expected

    @pytest.mark.parametrize("detail", ["Monitor", "Cena 14 Feb", "Algo 31/2/26"])
    def test_detail_without_valid_date(self, detail):
        assert parse_detail(detail) is None

    @pytest.mark.parametrize(
        "value, cents", [(1064.42, 106442), (0.1 + 0.2, 30), (1536.7333, 153673), (-683.62, -68362)]
    )
    def test_to_cents(self, value, cents):
        assert to_cents(value) == cents

    @pytest.mark.parametrize(
        "description, category",
        [
            ("UberEats", "food"),
            ("Uber", "transport"),
            ("Rappi", "food"),
            ("Gasolina", "fuel"),
            ("Interne", "utilities"),
            ("Steam", "entertainment"),
            ("Maestría", "education"),
            ("DisneyNewPlan", "subscriptions"),
            ("Amazon", "other"),
            ("Excel", "other"),  # "cel" solo cuenta al inicio de una palabra
        ],
    )
    def test_guess_category(self, description, category):
        assert guess_category(description) == category


class TestBuildBackup:
    def test_keeps_only_billed_current_and_active_installments(self):
        report = build([
            row("Viejo", date(2026, 8, 20), 10000),  # estado de septiembre: fuera
            row("Facturado", date(2026, 9, 5), 20000),  # día de corte → facturado
            row("Actual", date(2026, 10, 5), 30000),  # día de corte → actual
            row("MSI liquidado", date(2026, 3, 1), 60000, months=3),  # terminó en mayo: fuera
            row("MSI vivo", date(2026, 7, 1), 60000, months=6),  # julio..diciembre
        ])
        descriptions = [t["description"] for t in report.data["transactions"]]
        assert descriptions == [
            "MSI vivo",
            "Facturado",
            "Actual",
            "Saldo inicial: mensualidades MSI ya pagadas",
        ]
        assert report.billed_total_cents == 20000 + 10000
        assert report.current_total_cents == 30000 + 10000
        assert report.pending_installments_cents == 10000
        assert report.prepaid_installments_cents == 30000  # julio, agosto, septiembre

    def test_initial_balance_payment_is_dated_before_billed_period(self):
        report = build([row("MSI", date(2026, 7, 1), 60000, months=6)])
        payment = report.data["transactions"][-1]
        assert payment["type"] == "payment"
        assert payment["date"] == "2026-09-04"
        assert payment["amountCents"] == 30000

    def test_refunds(self):
        report = build([
            row("Compra", date(2026, 10, 1), 50000),
            row("Devolución vieja", date(2026, 8, 1), -1000),
            row("Devolución", date(2026, 10, 8), -2000),
        ])
        refunds = [t for t in report.data["transactions"] if t["type"] == "refund"]
        assert [(r["description"], r["amountCents"]) for r in refunds] == [("Devolución", 2000)]
        assert report.current_total_cents == -2000
        assert report.billed_total_cents == 50000

    def test_billed_paid_adds_payment_for_exact_statement_total(self):
        report = build([row("Compra", date(2026, 10, 1), 12345)], billed_paid_on=TODAY)
        payment = report.data["transactions"][-1]
        assert (payment["type"], payment["date"], payment["amountCents"]) == ("payment", "2026-10-23", 12345)

    def test_billed_payment_before_cut_is_rejected(self):
        with pytest.raises(ValueError, match="posterior al corte"):
            build([row("Compra", date(2026, 10, 1), 100)], billed_paid_on=date(2026, 10, 1))

    def test_warns_when_excel_first_installment_month_differs(self):
        mismatched = Row(7, "Gym", date(2026, 10, 1), 1000, 1, date(2026, 11, 1))
        report = build([mismatched])
        assert len(report.warnings) == 1
        assert "Fila 7" in report.warnings[0]

    def test_backup_shape(self):
        data = build([row("Compra", date(2026, 10, 1), 100, months=3)]).data
        assert data["schemaVersion"] == 2
        card = data["cards"][0]
        assert card == {
            "id": "id-1",
            "name": "BBVA Platinum",
            "cutDay": 5,
            "cutDayInNextPeriod": True,
            "dueDay": 25,
            "paymentBufferDays": 2,
            "createdAt": "2026-10-23T18:00:00.000Z",
            "updatedAt": "2026-10-23T18:00:00.000Z",
            "deletedAt": None,
        }
        assert data["settings"] == {"activeCardId": "id-1", "extraHolidays": []}
        assert len(data["categories"]) == 13
        tx = data["transactions"][0]
        assert tx["installments"] == {"months": 3}
        assert tx["cardId"] == "id-1"


def write_workbook(path, rows):
    workbook = Workbook()
    workbook.active.title = "- AYUDA -"
    sheet = workbook.create_sheet("Datos compra")
    sheet["B4"] = "Ingrese en la tabla los datos de compra"
    headers = ["TARJETA", "MONTO TOTAL", "# DE CUOTAS", "MONTO POR CUOTA",
               "MES DE 1ER CUOTA", "DETALLE", "MES ÚLTIMA CUOTA"]
    for col, header in enumerate(headers, start=2):
        sheet.cell(5, col, header)
    for r, (amount, months, first, detail) in enumerate(rows, start=6):
        sheet.cell(r, 2, "Visa")
        sheet.cell(r, 3, amount)
        sheet.cell(r, 4, months)
        sheet.cell(r, 6, first)
        sheet.cell(r, 7, detail)
    workbook.save(path)


class TestExcel:
    def test_read_rows(self, tmp_path):
        excel = tmp_path / "compras.xlsx"
        write_workbook(excel, [
            (1064.42, 3, datetime(2024, 2, 1), "Sin fecha"),
            (0, 1, datetime(2026, 11, 1), "Comida 20/10/26"),
            (500, 1, datetime(2026, 11, 1), "Gasolina 26/10/26"),
            (-683.62, 1, datetime(2026, 11, 1), "GlobalRefund 8/10/26"),
        ])
        result = read_rows(excel)
        assert result.without_date == [(6, "Sin fecha")]
        assert result.skipped_zero == [7]
        assert [(r.description, r.amount_cents, r.first_month) for r in result.rows] == [
            ("Gasolina", 50000, date(2026, 11, 1)),
            ("GlobalRefund", -68362, date(2026, 11, 1)),
        ]

    def test_main_writes_backup(self, tmp_path, capsys):
        excel = tmp_path / "compras.xlsx"
        write_workbook(excel, [(500, 1, datetime(2026, 11, 1), "Gasolina 6/10/26")])
        output = tmp_path / "salida.json"

        exit_code = main([
            str(excel), "--tarjeta", "BBVA Platinum", "--corte", "5", "--pago", "25",
            "--dia-de-corte-al-siguiente", "--hoy", "2026-10-23", "--salida", str(output),
        ])

        assert exit_code == 0
        data = json.loads(output.read_text(encoding="utf-8"))
        assert data["cards"][0]["cutDayInNextPeriod"] is True
        assert [t["description"] for t in data["transactions"]] == ["Gasolina"]
        assert "Periodo actual: $500.00" in capsys.readouterr().out
