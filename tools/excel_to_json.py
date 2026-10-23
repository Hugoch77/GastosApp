"""Convierte el Excel de compras a crédito en un respaldo JSON que GastosApp puede importar.

Uso (desde la raíz del proyecto):
    .venv/Scripts/python tools/excel_to_json.py RUTA_DEL_EXCEL \\
        --tarjeta "Mi tarjeta" --corte 5 --pago 25 --dia-de-corte-al-siguiente \\
        --facturado-pagado

Formato esperado del Excel (hoja "Datos compra"):
    - Una fila de encabezados con: MONTO TOTAL, # DE CUOTAS, MES DE 1ER CUOTA, DETALLE.
    - DETALLE termina con la fecha de compra en formato dd/mm/aa, p. ej. "Tacos 8/10/26".
    - MONTO TOTAL negativo = devolución. Las filas con monto 0 se ignoran.

Qué se importa (calculado con las mismas reglas que la app, ver src/domain/billing.ts):
    - Movimientos del estado facturado y del periodo actual.
    - Compras a meses sin intereses con mensualidades pendientes, con su fecha original.
    - Un pago "saldo inicial" por las mensualidades MSI que ya se pagaron antes del estado
      facturado, para que el saldo total cuadre.
    - Con --facturado-pagado, un pago por el total del estado facturado.

El archivo resultante contiene datos personales: se guarda junto al Excel, fuera del repositorio.
"""

import argparse
import calendar
import json
import re
import sys
import unicodedata
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import load_workbook

SHEET_NAME = "Datos compra"
SCHEMA_VERSION = 2
# México no tiene horario de verano desde 2022: la Ciudad de México es UTC-6 todo el año.
MEXICO_CITY = timezone(timedelta(hours=-6))
DATE_IN_DETAIL = re.compile(r"\s*(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})\s*$")

HEADERS = {
    "amount": "MONTO TOTAL",
    "months": "# DE CUOTAS",
    "first_month": "MES DE 1ER CUOTA",
    "detail": "DETALLE",
}

# Deben coincidir con DEFAULT_CATEGORIES de src/data/defaults.ts.
DEFAULT_CATEGORIES = [
    ("food", "Comida", "🍔"),
    ("groceries", "Supermercado", "🛒"),
    ("transport", "Transporte", "🚗"),
    ("fuel", "Gasolina", "⛽"),
    ("utilities", "Servicios", "💡"),
    ("health", "Salud", "💊"),
    ("entertainment", "Entretenimiento", "🎬"),
    ("clothing", "Ropa", "👕"),
    ("home", "Hogar", "🏠"),
    ("education", "Educación", "📚"),
    ("travel", "Viajes", "✈️"),
    ("subscriptions", "Suscripciones", "📺"),
    ("other", "Otros", "📦"),
]

# Se revisan en orden: la primera palabra clave que aparezca decide la categoría.
# Cada palabra clave debe coincidir con el inicio de una palabra ("interne" → "Internet").
CATEGORY_KEYWORDS = [
    ("food", ["ubereats", "uber eats", "rappi", "didi food", "tacos", "comida", "pizza",
              "restaurante", "cafe", "birria", "cena", "desayuno", "sushi", "cerveza"]),
    ("subscriptions", ["netflix", "spotify", "disney", "hbo", "prime video", "game pass",
                       "youtube", "icloud", "vshred", "antivirus"]),
    ("fuel", ["gasolina"]),
    ("transport", ["uber", "didi", "dott", "lime", "taxi", "caseta", "estacionamiento",
                   "placas", "multa", "llantas", "carro", "bus"]),
    ("groceries", ["walmart", "soriana", "chedraui", "costco", "sams", "oxxo", "super",
                   "bodega aurrera", "heb"]),
    ("utilities", ["cfe", "luz", "agua", "interne", "telmex", "izzi", "totalplay", "cel",
                   "telefono", "at&t", "telcel", "recarga"]),
    ("health", ["farmacia", "medic", "dentalia", "dentista", "dientes", "doctor", "hospital", "gym"]),
    ("entertainment", ["steam", "juego", "xbox", "playstation", "nintendo", "cine", "tiktok", "wow"]),
    ("clothing", ["shein", "ropa", "zapatos", "tenis"]),
    ("home", ["lavadora", "roomba", "caminadora", "ferreteria", "home depot", "cafetera"]),
    ("education", ["maestria", "curso", "udemy", "libro"]),
    ("travel", ["volaris", "viva aerobus", "vivaaerobus", "aeromexico", "airbnb", "hotel",
                "vuelo"]),
]


# ---------------------------------------------------------------------------
# Reglas de fechas (mismas que src/domain/billing.ts)
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class CardRules:
    cut_day: int
    due_day: int
    cut_day_in_next_period: bool


def clamped_date(year: int, month: int, day: int) -> date:
    """Un día mayor a los días del mes se recorta al último día (31 → 28 en febrero)."""
    return date(year, month, min(day, calendar.monthrange(year, month)[1]))


def shift_month(year: int, month: int, months: int) -> tuple[int, int]:
    index = year * 12 + (month - 1) + months
    return index // 12, index % 12 + 1


def cut_relative_to(d: date, cut_day: int, months: int) -> date:
    year, month = shift_month(d.year, d.month, months)
    return clamped_date(year, month, cut_day)


def statement_cut_for(d: date, rules: CardRules) -> date:
    """Fecha de corte del estado de cuenta al que pertenece un movimiento del día `d`."""
    if rules.cut_day_in_next_period:
        d = d + timedelta(days=1)
    cut = cut_relative_to(d, rules.cut_day, 0)
    return cut if cut >= d else cut_relative_to(d, rules.cut_day, 1)


def period_for_cut(cut: date, rules: CardRules) -> tuple[date, date]:
    """Fechas (inclusive) del periodo que cierra en `cut`."""
    previous_cut = cut_relative_to(cut, rules.cut_day, -1)
    if rules.cut_day_in_next_period:
        return previous_cut, cut - timedelta(days=1)
    return previous_cut + timedelta(days=1), cut


def due_date_for_cut(cut: date, due_day: int) -> date:
    same_month = clamped_date(cut.year, cut.month, due_day)
    if same_month > cut:
        return same_month
    year, month = shift_month(cut.year, cut.month, 1)
    return clamped_date(year, month, due_day)


def split_installments(total_cents: int, months: int) -> list[int]:
    """floor(total / meses); el residuo, de 1 en 1 centavo en las primeras cuotas."""
    base, remainder = divmod(total_cents, months)
    return [base + 1 if i < remainder else base for i in range(months)]


def installment_schedule(
    purchase: date, total_cents: int, months: int, rules: CardRules
) -> list[tuple[date, int]]:
    """(fecha de corte, monto) de cada mensualidad."""
    cut = statement_cut_for(purchase, rules)
    schedule = []
    for amount in split_installments(total_cents, months):
        schedule.append((cut, amount))
        cut = cut_relative_to(cut, rules.cut_day, 1)
    return schedule


# ---------------------------------------------------------------------------
# Lectura del Excel
# ---------------------------------------------------------------------------


@dataclass
class Row:
    number: int
    description: str
    purchase_date: date
    amount_cents: int
    months: int
    first_month: date | None


@dataclass
class ReadResult:
    rows: list[Row] = field(default_factory=list)
    skipped_zero: list[int] = field(default_factory=list)
    without_date: list[tuple[int, str]] = field(default_factory=list)


def to_cents(value: object) -> int:
    return int((Decimal(str(value)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def parse_detail(detail: str) -> tuple[str, date] | None:
    """"Tacos 8/10/26" → ("Tacos", 2026-10-08). None si no termina en una fecha válida."""
    match = DATE_IN_DETAIL.search(detail)
    if not match:
        return None
    day, month, year = (int(g) for g in match.groups())
    if year < 100:
        year += 2000
    try:
        purchase = date(year, month, day)
    except ValueError:
        return None
    description = detail[: match.start()].strip() or "Compra"
    return description[:80], purchase


def find_header_row(sheet) -> tuple[int, dict[str, int]]:
    for row in sheet.iter_rows(min_row=1, max_row=50):
        texts = {str(c.value).strip().upper(): c.column for c in row if c.value is not None}
        if all(h in texts for h in HEADERS.values()):
            return row[0].row, {key: texts[h] for key, h in HEADERS.items()}
    raise ValueError(f"No se encontró la fila de encabezados {list(HEADERS.values())}")


def read_rows(excel_path: Path) -> ReadResult:
    workbook = load_workbook(excel_path, data_only=True, read_only=False)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f'El Excel no tiene la hoja "{SHEET_NAME}"')
    sheet = workbook[SHEET_NAME]
    header_row, columns = find_header_row(sheet)

    result = ReadResult()
    for r in range(header_row + 1, sheet.max_row + 1):
        def cell(key: str, row: int = r) -> object:
            return sheet.cell(row, columns[key]).value

        amount, detail = cell("amount"), cell("detail")
        if amount is None and detail is None:
            continue
        amount_cents = to_cents(amount or 0)
        if amount_cents == 0:
            result.skipped_zero.append(r)
            continue
        parsed = parse_detail(str(detail or ""))
        if parsed is None:
            result.without_date.append((r, str(detail)))
            continue
        description, purchase = parsed
        first_month = cell("first_month")
        result.rows.append(
            Row(
                number=r,
                description=description,
                purchase_date=purchase,
                amount_cents=amount_cents,
                months=int(cell("months") or 1),
                first_month=first_month.date() if isinstance(first_month, datetime) else None,
            )
        )
    return result


def guess_category(description: str) -> str:
    text = unicodedata.normalize("NFKD", description.lower())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    for category, keywords in CATEGORY_KEYWORDS:
        if any(re.search(r"(?<![a-z0-9])" + re.escape(k), text) for k in keywords):
            return category
    return "other"


# ---------------------------------------------------------------------------
# Construcción del respaldo
# ---------------------------------------------------------------------------


@dataclass
class ImportReport:
    data: dict
    billed_cut: date
    billed_total_cents: int
    current_total_cents: int
    pending_installments_cents: int
    prepaid_installments_cents: int
    imported_rows: int
    warnings: list[str]


def timestamp(now: datetime) -> str:
    return now.astimezone(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def build_backup(
    rows: list[Row],
    *,
    card_name: str,
    rules: CardRules,
    payment_buffer_days: int,
    today: date,
    billed_paid_on: date | None,
    now: datetime,
    new_id=lambda: str(uuid.uuid4()),
) -> ImportReport:
    current_cut = statement_cut_for(today, rules)
    billed_cut = cut_relative_to(current_cut, rules.cut_day, -1)
    billed_start, _ = period_for_cut(billed_cut, rules)
    current_start, _ = period_for_cut(current_cut, rules)
    created = timestamp(now)
    card_id = new_id()
    warnings: list[str] = []

    transactions: list[dict] = []
    totals = {"billed": 0, "current": 0, "pending": 0, "prepaid": 0}

    def add(tx_type: str, tx_date: date, description: str, amount_cents: int,
            category: str | None, months: int | None) -> None:
        transactions.append({
            "id": new_id(),
            "cardId": card_id,
            "date": tx_date.isoformat(),
            "description": description,
            "amountCents": amount_cents,
            "type": tx_type,
            "categoryId": category,
            "installments": {"months": months} if months else None,
            "createdAt": created,
            "updatedAt": created,
            "deletedAt": None,
        })

    imported = 0
    for row in sorted(rows, key=lambda r: (r.purchase_date, r.number)):
        if row.amount_cents < 0:
            if row.purchase_date < billed_start:
                continue
            refund = -row.amount_cents
            totals["billed" if row.purchase_date < current_start else "current"] -= refund
            add("refund", row.purchase_date, row.description, refund, guess_category(row.description), None)
            imported += 1
            continue

        if row.months > 48:
            warnings.append(f"Fila {row.number}: {row.months} meses no es válido (máximo 48); se omitió")
            continue
        schedule = installment_schedule(row.purchase_date, row.amount_cents, row.months, rules)
        if schedule[-1][0] < billed_cut:
            continue  # ya liquidada antes del estado facturado

        first_due = due_date_for_cut(schedule[0][0], rules.due_day)
        if row.first_month and (row.first_month.year, row.first_month.month) != (first_due.year, first_due.month):
            warnings.append(
                f"Fila {row.number} ({row.description} {row.purchase_date:%d/%m/%y}): en el Excel "
                f"la primera cuota es {row.first_month:%m/%Y}, la app la calcula en {first_due:%m/%Y}"
            )

        for cut, amount in schedule:
            if cut < billed_cut:
                totals["prepaid"] += amount
            elif cut == billed_cut:
                totals["billed"] += amount
            elif cut == current_cut:
                totals["current"] += amount
            else:
                totals["pending"] += amount

        add("charge", row.purchase_date, row.description, row.amount_cents,
            guess_category(row.description), row.months if row.months > 1 else None)
        imported += 1

    if totals["prepaid"] > 0:
        # Antes del periodo facturado: no afecta el pago para no generar intereses.
        add("payment", billed_start - timedelta(days=1),
            "Saldo inicial: mensualidades MSI ya pagadas", totals["prepaid"], None, None)

    if billed_paid_on is not None and totals["billed"] > 0:
        if billed_paid_on < current_start:
            raise ValueError(
                f"La fecha del pago ({billed_paid_on}) debe ser posterior al corte del {billed_cut}"
            )
        add("payment", billed_paid_on, "Pago a tarjeta", totals["billed"], None, None)

    data = {
        "schemaVersion": SCHEMA_VERSION,
        "settings": {"activeCardId": card_id, "extraHolidays": []},
        "cards": [{
            "id": card_id,
            "name": card_name,
            "cutDay": rules.cut_day,
            "cutDayInNextPeriod": rules.cut_day_in_next_period,
            "dueDay": rules.due_day,
            "paymentBufferDays": payment_buffer_days,
            "createdAt": created,
            "updatedAt": created,
            "deletedAt": None,
        }],
        "categories": [
            {"id": cid, "name": name, "icon": icon, "updatedAt": created, "deletedAt": None}
            for cid, name, icon in DEFAULT_CATEGORIES
        ],
        "transactions": transactions,
        "exportedAt": created,
    }
    return ImportReport(
        data=data,
        billed_cut=billed_cut,
        billed_total_cents=totals["billed"],
        current_total_cents=totals["current"],
        pending_installments_cents=totals["pending"],
        prepaid_installments_cents=totals["prepaid"],
        imported_rows=imported,
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# Línea de comandos
# ---------------------------------------------------------------------------


def money(cents: int) -> str:
    sign = "-" if cents < 0 else ""
    return f"{sign}${abs(cents) / 100:,.2f}"


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Excel de compras → respaldo JSON de GastosApp")
    parser.add_argument("excel", type=Path, help="ruta del archivo .xlsx")
    parser.add_argument("--tarjeta", required=True, help='nombre de la tarjeta, p. ej. "BBVA Platinum"')
    parser.add_argument("--corte", type=int, required=True, help="día de corte (1-31)")
    parser.add_argument("--pago", type=int, required=True, help="día límite de pago (1-31)")
    parser.add_argument("--dias-antes", type=int, default=2,
                        help="días antes de la fecha límite en que pagas (default 2)")
    parser.add_argument("--dia-de-corte-al-siguiente", action="store_true",
                        help="las compras del día de corte entran al estado siguiente")
    parser.add_argument("--facturado-pagado", nargs="?", const="hoy", metavar="AAAA-MM-DD",
                        help="registra el pago del estado facturado (fecha opcional; default hoy)")
    parser.add_argument("--hoy", type=date.fromisoformat,
                        help="fecha de referencia AAAA-MM-DD (default: hoy en la Ciudad de México)")
    parser.add_argument("--salida", type=Path, help="archivo JSON de salida (default: junto al Excel)")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    # La consola de Windows no siempre usa UTF-8; así los acentos no rompen la salida.
    sys.stdout.reconfigure(encoding="utf-8")
    args = parse_args(sys.argv[1:] if argv is None else argv)
    for name, day in (("--corte", args.corte), ("--pago", args.pago)):
        if not 1 <= day <= 31:
            print(f"{name} debe estar entre 1 y 31", file=sys.stderr)
            return 2

    now = datetime.now(MEXICO_CITY)
    today = args.hoy or now.date()
    paid_on = None
    if args.facturado_pagado:
        paid_on = today if args.facturado_pagado == "hoy" else date.fromisoformat(args.facturado_pagado)

    read = read_rows(args.excel)
    report = build_backup(
        read.rows,
        card_name=args.tarjeta,
        rules=CardRules(args.corte, args.pago, args.dia_de_corte_al_siguiente),
        payment_buffer_days=args.dias_antes,
        today=today,
        billed_paid_on=paid_on,
        now=now,
    )

    output = args.salida or args.excel.with_name(f"gastosapp-importacion-{today.isoformat()}.json")
    output.write_text(json.dumps(report.data, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Filas leídas: {len(read.rows)} · importadas: {report.imported_rows}")
    print(f"Estado facturado (corte {report.billed_cut:%d/%m/%Y}): {money(report.billed_total_cents)}"
          + (" · pagado" if paid_on else ""))
    print(f"Periodo actual: {money(report.current_total_cents)}")
    print(f"MSI por facturar: {money(report.pending_installments_cents)}")
    if report.prepaid_installments_cents:
        print(f"Mensualidades MSI ya pagadas (saldo inicial): {money(report.prepaid_installments_cents)}")
    if read.skipped_zero:
        print(f"Filas con monto 0 ignoradas: {', '.join(map(str, read.skipped_zero))}")
    if read.without_date:
        print(f"Filas sin fecha dd/mm/aa en DETALLE (ignoradas): {len(read.without_date)}")
    for warning in report.warnings:
        print(f"Aviso: {warning}")
    print(f"\nListo: {output}")
    print("Impórtalo en la app desde Ajustes → Importar respaldo (o en la pantalla de bienvenida).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
